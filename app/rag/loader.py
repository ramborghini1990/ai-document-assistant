import io
import os
from typing import List, Dict, Any
from pypdf import PdfReader
from docx import Document as DocxDocument
from PIL import Image
from app.ai.gemini import get_gemini_client, MODEL_NAME
from google.genai import types


def extract_text_from_pdf(file_stream, filename: str) -> List[Dict[str, Any]]:
    """استخراج متن از صفحات فایل PDF متنی."""
    if not file_stream:
        raise ValueError(f"File stream for '{filename}' is empty or invalid.")

    try:
        reader = PdfReader(file_stream)
    except Exception as e:
        raise ValueError(f"Failed to read PDF file '{filename}': {str(e)}")

    if not reader.pages:
        raise ValueError(f"The PDF file '{filename}' contains no pages.")

    pages_data = []
    for idx, page in enumerate(reader.pages):
        page_num = idx + 1
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        cleaned_text = " ".join(text.split())
        if cleaned_text:
            pages_data.append({
                "page_number": page_num,
                "text": cleaned_text
            })

    return pages_data


def extract_text_from_docx(file_stream, filename: str) -> List[Dict[str, Any]]:
    """استخراج ساختاریافته متن و جداول از فایل Word (.docx)."""
    if not file_stream:
        raise ValueError(f"File stream for '{filename}' is empty or invalid.")

    try:
        doc = DocxDocument(file_stream)
    except Exception as e:
        raise ValueError(f"Failed to read Word document '{filename}': {str(e)}")

    content_blocks = []

    # ۱. استخراج پاراگراف‌ها و تیترها
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_blocks.append(text)

    # ۲. استخراج جداول (سطر به سطر با تفکیک ستون‌ها)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # حذف مقادیر تکراری سلول‌های ادغام‌شده (merged cells)
                unique_cells = []
                for c in row_cells:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)
                content_blocks.append(" | ".join(unique_cells))

    full_text = "\n".join(content_blocks).strip()
    if not full_text:
        return []

    # برای اسناد ورد کل محتوا به عنوان بخش ۱ منظور می‌شود
    return [{"page_number": 1, "text": full_text}]


def extract_text_from_image(file_stream, filename: str) -> List[Dict[str, Any]]:
    """استخراج متن از تصویر با فشرده‌سازی ابعاد و کیفیت جهت مصرف حداقل حجم و پهنای باند."""
    if not file_stream:
        raise ValueError(f"File stream for '{filename}' is empty or invalid.")

    try:
        file_stream.seek(0)
        file_bytes = file_stream.read()
        image = Image.open(io.BytesIO(file_bytes))

        # تبدیل به RGB برای خلاص شدن از کانال آلفا و کاهش حجم
        if image.mode != "RGB":
            image = image.convert("RGB")

        # تغییر ابعاد به حداکثر 1280 پیکسل (کاملاً خوانا برای OCR اما بسیار سبک)
        max_size = 1280
        if max(image.size) > max_size:
            image.thumbnail((max_size, max_size), Image.Resampling.LANCZOS)

        # فشرده‌سازی در قالب بایت‌های JPEG با کیفیت 75%
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=75, optimize=True)
        buffer.seek(0)
        compressed_image = Image.open(buffer)

    except Exception as e:
        raise ValueError(f"Failed to process/compress image '{filename}': {str(e)}")

    vision_prompt = (
        "You are an expert document OCR assistant. "
        "Extract all readable text, titles, dates, numbers, and tabular data verbatim. "
        "Do not summarize. Return pure transcribed text."
    )

    from app.ai.gemini import transcribe_image_with_vision
    transcribed_text = transcribe_image_with_vision(compressed_image, vision_prompt)

    if not transcribed_text:
        return []

    return [{"page_number": 1, "text": transcribed_text}]


def load_document_content(file, filename: str) -> List[Dict[str, Any]]:
    """توزیع‌کننده یکپارچه (Unified Dispatcher) جهت انتخاب پردازشگر مناسب بر اساس پسوند فایل."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file, filename)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file, filename)
    elif ext in [".jpg", ".jpeg", ".png", ".webp"]:
        return extract_text_from_image(file, filename)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Supported: PDF, DOCX, JPG, PNG.")