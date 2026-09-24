import io
from datetime import datetime
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# نگاشت فیلدهای فنی به ستون‌های رسمی سازمانی (مطابق با مرجع ISO 14001)
COLUMN_LABELS = {
    # Scadenziario
    "tema": "Tema / Area",
    "adempimento": "Adempimento / Descrizione",
    "scadenza": "Data Scadenza",
    "preavviso": "Preavviso / Frequenza",
    "responsabile": "Responsabile",
    "stato": "Stato",
    "evidenza_note": "Evidenza / Protocollo / Note",
    "allerta": "Allerta / Priorità",
    
    # Vehicle & Fuel
    "targa": "Targa Mezzo",
    "modello": "Modello / Descrizione",
    "scadenza_revisione": "Scadenza Revisione",
    "periodo": "Periodo Riferimento",
    "litri": "Litri Carburante",
    "chilometri": "Km Percorsi",
    "consumo_l_100km": "Consumo (L/100km)",
    
    # General
    "titolo_documento": "Titolo Documento",
    "ente_emittente": "Ente Emittente",
    "numero_protocollo": "Numero Protocollo",
    "data_emissione": "Data Emissione",
    "data_scadenza": "Data Scadenza",
    "oggetto_sintesi": "Oggetto Sintesi",
    
    # Metadata
    "source_page": "Pagina Fonte",
    "source_document": "Documento Fonte"
}


def create_styled_excel_report(
    extracted_data_by_schema: Dict[str, Any],
    source_documents: List[str]
) -> bytes:
    """
    تولید فایل اکسل تجاری و چندشیت مطابق استانداردهای ممیزی و ایزو ۱۴۰۰۱.
    """
    wb = openpyxl.Workbook()
    # حذف شیت پیش‌فرض خالی
    wb.remove(wb.active)

    # استایل‌های سازمانی ISO
    header_fill = PatternFill(start_color="1E4D2B", end_color="1E4D2B", fill_type="solid")  # سبز زیتونی سازمانی
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    meta_font = Font(name="Segoe UI", size=10, italic=True, color="555555")
    calculated_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")  # سبز بسیار ملایم برای مقادیر محاسباتی
    
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    # -------------------------------------------------------------
    # شیت ۱: خلاصه مدیریتی (Riepilogo / Summary)
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="RIEPILOGO_GENERALE")
    ws_summary.views.sheetView[0].showGridLines = True

    # تیتر گزارش
    ws_summary["A1"] = "REPORT DOCUMENT INTELLIGENCE & COMPLIANCE"
    ws_summary["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E4D2B")
    
    ws_summary["A2"] = f"Generato il: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws_summary["A2"].font = meta_font

    # جدول آمار اسناد
    ws_summary["A4"] = "Documenti Sorgente Inclusi:"
    ws_summary["A4"].font = Font(name="Segoe UI", size=11, bold=True)
    
    row_idx = 5
    for doc in source_documents:
        ws_summary[f"A{row_idx}"] = f"• {doc}"
        ws_summary[f"A{row_idx}"].font = data_font
        row_idx += 1

    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Riepilogo Schemi Estratti:"
    ws_summary[f"A{row_idx}"].font = Font(name="Segoe UI", size=11, bold=True)
    row_idx += 1

    # هدر جدول خلاصه
    sum_headers = ["Schema Dati", "Totale Record", "Stato Estrazione", "Foglio Dedicato"]
    for col_i, h_text in enumerate(sum_headers, start=1):
        cell = ws_summary.cell(row=row_idx, column=col_i, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx += 1
    for schema_name, data in extracted_data_by_schema.items():
        recs = data.get("records", [])
        ws_summary.cell(row=row_idx, column=1, value=schema_name.upper()).font = data_font
        ws_summary.cell(row=row_idx, column=2, value=len(recs)).font = data_font
        ws_summary.cell(row=row_idx, column=3, value="Completato").font = data_font
        ws_summary.cell(row=row_idx, column=4, value=schema_name.upper()).font = data_font
        row_idx += 1

    # -------------------------------------------------------------
    # شیت‌های داده‌های تخصصی (Data Sheets)
    # -------------------------------------------------------------
    for schema_name, data in extracted_data_by_schema.items():
        sheet_title = schema_name.upper()[:31]
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"  # فریز کردن سطر عنوان

        records = data.get("records", [])
        if not records:
            ws.cell(row=1, column=1, value="Nessun dato estratto per questo schema.").font = meta_font
            continue

        # مشخص کردن لیست ستون‌ها
        sample_fields = list(records[0].get("fields", {}).keys())
        headers = ["Pagina Fonte"] + [COLUMN_LABELS.get(f, f.replace("_", " ").title()) for f in sample_fields] + ["Stato Estrazione"]

        # ایجاد ردیف هدر
        for col_i, h_name in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_i, value=h_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # درج سطرهای داده
        for r_idx, record in enumerate(records, start=2):
            src_page = record.get("source_page", 1)
            fields = record.get("fields", {})

            # ستون شماره صفحه
            p_cell = ws.cell(row=r_idx, column=1, value=f"Pagina {src_page}")
            p_cell.font = meta_font
            p_cell.alignment = Alignment(horizontal="center")
            p_cell.border = thin_border

            # درج فیلدها
            overall_status = "EXTRACTED"
            for c_idx, fname in enumerate(sample_fields, start=2):
                fdata = fields.get(fname, {})
                norm_val = fdata.get("normalized_value")
                raw_val = fdata.get("raw_value")
                f_status = fdata.get("status", "EXTRACTED")
                if f_status == "CALCULATED":
                    overall_status = "CALCULATED"

                val_to_write = norm_val if norm_val is not None else raw_val
                cell = ws.cell(row=r_idx, column=c_idx, value=val_to_write)
                cell.font = data_font
                cell.border = thin_border

                if f_status == "CALCULATED":
                    cell.fill = calculated_fill

                # تراز وسط برای تاریخ‌ها و مقادیر کوتاه
                if any(k in fname for k in ["data", "scadenza", "targa", "periodo", "consumo"]):
                    cell.alignment = Alignment(horizontal="center")

            # ستون وضعیت ردیف
            st_cell = ws.cell(row=r_idx, column=len(headers), value=overall_status)
            st_cell.font = Font(name="Segoe UI", size=9, bold=True, color="2E7D32" if overall_status == "EXTRACTED" else "1565C0")
            st_cell.alignment = Alignment(horizontal="center")
            st_cell.border = thin_border

    # تنظیم خودکار پهنای ستون‌ها برای تمام شیت‌ها
    for ws_item in wb.worksheets:
        for col in ws_item.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_item.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()