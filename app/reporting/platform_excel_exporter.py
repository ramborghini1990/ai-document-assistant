"""
Platform Multi-Sheet Enterprise Excel Exporter (v2).

Sheet 1: Dashboard & KPI
Sheet 2: Registro Documenti   (one row per uploaded file, metadata derived from SCHEMA definitions)
Sheet 3: Scadenze & Azioni    (deadlines derived from the same in-memory records, not from stale DB rows)
Sheet 4+: Dati <schema>       (analytical sheet for every schema that has extracted records)

Design rules (fixes vs v1):
- Field mapping is driven by exact schema field names and by `alert` definitions in schemas/*.json,
  not by fragile substring matching.
- Records of files that are no longer uploaded are excluded everywhere.
- A document with several records gets the EARLIEST due date and the DISTINCT subjects/issuers.
- Dates are written as real Excel dates; text starting with '=' is never interpreted as a formula.
- No hard-coded placeholder texts ("Verificato e Registrato", "Responsabile Documentale").
"""
import io
import json
import math
from datetime import date, datetime, timedelta
from typing import Any, Dict, Iterator, List, Optional, Tuple

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.alerts.deadlines import extract_deadlines_from_records
from app.extraction.registry import SchemaDef, get_module, resolve_schema
from app.reporting.analytics import ConsumptionAnalysis, MONTH_FULL, analyze_payload, period_label
from app.reporting.charts import render_monthly_chart
from app.reporting.common import cell_value, column_plan, sheet_title_for
from app.reporting.expansion import expand_record, unit_key

OK_STATUSES = ("EXTRACTED", "EDITED", "CALCULATED")
MANUAL_SOURCES = ("", "manual", None)

# Exact field names declared in schemas/*.json, in priority order.
SUBJECT_KEYS = (
    "nominativo_dipendente", "macchinario_asset", "targa", "codice_pod", "codice_pod_pdr",
    "adempimento", "titolo_documento", "numero_atto_protocollo", "tipo_utenza",
)
ISSUER_KEYS = ("ente_formatore", "ente_emittente", "ente_rilascio", "fornitore", "fornitore_tecnico")
ISSUE_DATE_KEYS = ("data_emissione", "data_rilascio", "data_intervento", "data_rifornimento")
RESPONSIBLE_KEYS = ("responsabile",)
COMPANY_KEYS = ("ragione_sociale",)
ADDRESS_KEYS = ("sede_indirizzo",)

GREEN = "1E4D2B"
DATE_FMT = "DD/MM/YYYY"

# Readability: sizes in points, row heights in points. Data rows never go below ROW_HEIGHT; wrapped text grows them.
FONT_DATA = 10
FONT_HEADER = 11
FONT_KPI_LABEL = 10
FONT_KPI_VALUE = 18
ROW_HEIGHT = 26
HEADER_HEIGHT = 34
LINE_HEIGHT = 14
COL_MIN_WIDTH, COL_MAX_WIDTH, COL_PADDING = 16, 60, 6
STATUS_COLOR = {"Scaduto": "C62828", "In scadenza": "EF6C00", "Valido": "2E7D32"}


# ----------------------------------------------------------------------------- helpers
def _font(**kw) -> Font:
    return Font(name="Segoe UI", **kw)


def _border() -> Border:
    thin = Side(style="thin", color="D0D5DD")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _to_date(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and len(v.strip()) == 10:
        try:
            return date.fromisoformat(v.strip())
        except ValueError:
            return None
    return None


def _calc_doc_status(due: Any, today: Optional[date] = None) -> Tuple[str, Optional[int]]:
    d = _to_date(due)
    if d is None:
        return "Senza scadenza", None
    days = (d - (today or date.today())).days
    if days < 0:
        return "Scaduto", days
    if days <= 30:
        return "In scadenza", days
    return "Valido", days


def _text(fdata: Any) -> Optional[str]:
    if not isinstance(fdata, dict) or fdata.get("status") == "MISSING":
        return None
    v = cell_value(fdata)
    s = str(v).strip() if v is not None else ""
    return s or None


def _date_of(fdata: Any) -> Optional[date]:
    if not isinstance(fdata, dict) or fdata.get("status", "EXTRACTED") not in OK_STATUSES:
        return None
    return _to_date(fdata.get("normalized_value"))


def _distinct(values: List[str], limit: int = 3) -> str:
    seen: List[str] = []
    for v in values:
        if v and v not in seen:
            seen.append(v)
    if not seen:
        return "-"
    extra = len(seen) - limit
    return "; ".join(seen[:limit]) + (f" (+{extra})" if extra > 0 else "")


def _schema_for(mid: str, sname: str) -> Optional[SchemaDef]:
    try:
        return resolve_schema(sname, mid)
    except Exception:
        return None


def _module_for(mid: str):
    try:
        return get_module(mid)
    except Exception:
        return None


def _mod_short(mid: str) -> str:
    m = _module_for(mid)
    return m.label.split("—")[0].strip() if m else mid


def _due_field_names(schema: Optional[SchemaDef], rec: Dict[str, Any]) -> List[str]:
    if schema:
        return [n for n, f in schema.fields.items() if f.alert]
    return [n for n in (rec.get("fields") or {}) if "scadenza" in n.lower()]


def _is_kept(src: Any, active: set) -> bool:
    return src in active or src in MANUAL_SOURCES


def _iter_records(payload: Dict[str, Any]) -> Iterator[Tuple[str, str, Dict[str, Any]]]:
    for mid, schemas in (payload or {}).items():
        for sname, sdata in (schemas or {}).items():
            for rec in ((sdata or {}).get("records") or []):
                yield mid, sname, rec


def _filter_payload(payload: Dict[str, Any], active: set) -> Dict[str, Dict[str, Any]]:
    out: Dict[str, Dict[str, Any]] = {}
    for mid, schemas in (payload or {}).items():
        for sname, sdata in (schemas or {}).items():
            recs = [r for r in ((sdata or {}).get("records") or []) if _is_kept(r.get("source_document"), active)]
            if recs:
                out.setdefault(mid, {})[sname] = {**sdata, "records": recs, "total_records": len(recs)}
    return out


def _cell_alignment(value: Any, align: Optional[Alignment]) -> Alignment:
    """Every cell is vertically centred; text wraps; dates and numbers are centred. An explicit `align` wins field by field."""
    is_text = isinstance(value, str)
    centred = isinstance(value, (date, datetime, int, float)) and not isinstance(value, bool)
    horizontal = align.horizontal if align and align.horizontal else ("center" if centred else None)
    vertical = align.vertical if align and align.vertical else "center"
    wrap = align.wrap_text if align and align.wrap_text is not None else is_text
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def _put(ws, row: int, col: int, value: Any, font: Optional[Font] = None,
         align: Optional[Alignment] = None, border: Optional[Border] = None, fill=None):
    """Safe cell writer: dates get a real date format, '=...' text is never a formula."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str) and value.startswith("="):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    if isinstance(value, (date, datetime)):
        cell.number_format = DATE_FMT
    if font:
        cell.font = font
    cell.alignment = _cell_alignment(value, align)
    if border:
        cell.border = border
    if fill:
        cell.fill = fill
    return cell


def _header(ws, headers: List[str], row: int = 1) -> None:
    fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
    for c, h in enumerate(headers, start=1):
        _put(ws, row, c, h, _font(size=FONT_HEADER, bold=True, color="FFFFFF"),
             Alignment(horizontal="center", vertical="center", wrap_text=True), fill=fill)
    ws.row_dimensions[row].height = HEADER_HEIGHT


def _unique_title(wb, base: str) -> str:
    base = base[:31]
    if base not in wb.sheetnames:
        return base
    i = 2
    while True:
        suffix = f" {i}"
        cand = base[: 31 - len(suffix)] + suffix
        if cand not in wb.sheetnames:
            return cand
        i += 1


def _longest_line(value: Any) -> int:
    return max((len(part) for part in str(value).split(chr(10))), default=0)


def _autosize(ws, skip_rows: int = 0) -> None:
    """Column width from the longest line of each column (plus padding); very long text wraps instead of widening."""
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((_longest_line(c.value) if c.value is not None and c.row > skip_rows else 0) for c in col)
        ws.column_dimensions[letter].width = min(max(longest + COL_PADDING, COL_MIN_WIDTH), COL_MAX_WIDTH)


def _fit_row_heights(ws, first_row: int = 2) -> None:
    """Data rows get at least ROW_HEIGHT; rows with wrapped text grow to the number of lines it needs (call after _autosize)."""
    for row in ws.iter_rows(min_row=first_row):
        lines = 1
        for cell in row:
            if isinstance(cell.value, str) and cell.alignment.wrap_text:
                width = ws.column_dimensions[get_column_letter(cell.column)].width or COL_MIN_WIDTH
                per_line = max(1, int(width) - 2)
                lines = max(lines, sum(max(1, math.ceil(len(part) / per_line)) for part in cell.value.split(chr(10))))
        ws.row_dimensions[row[0].row].height = min(409, max(ROW_HEIGHT, lines * LINE_HEIGHT + 8))


def _autofilter(ws) -> None:
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _record_status(rec: Dict[str, Any]) -> str:
    if rec.get("warnings"):
        return "Da verificare"
    sts = {f.get("status") for f in (rec.get("fields") or {}).values() if isinstance(f, dict)}
    return "Modificato" if "EDITED" in sts else "Estratto"


# ----------------------------------------------------------------------------- consumption trend sheet
TREND_SHEET = "Trend & Analisi Consumi"
CHART_WIDTH_PX = 720
NUM_FMT = "#,##0.00"


def write_trend_sheet(wb, analyses: List[ConsumptionAnalysis]) -> None:
    """Monthly matrix (months x years, totals, averages, like-for-like YoY) + chart next to it, for every analysed schema."""
    ws = wb.create_sheet(TREND_SHEET)
    border = _border()
    data_font = _font(size=FONT_DATA)
    bold = _font(size=FONT_DATA, bold=True)
    meta_font = _font(size=FONT_DATA, italic=True, color="555555")
    total_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    _put(ws, 1, 1, "TREND & ANALISI CONSUMI", _font(size=16, bold=True, color=GREEN))
    _put(ws, 2, 1, "Somme dei soli valori estratti (nessuna stima); i periodi plurimensili sono ripartiti in parti uguali: "
                   "valori ripartiti pro-quota sui mesi di competenza. Mesi senza dato = '-'. Variazioni calcolate sui soli "
                   "mesi presenti in entrambi gli anni.", meta_font, Alignment(wrap_text=False))
    row = 4
    max_cols = 1
    for a in analyses:
        _put(ws, row, 1, f"{a.label} — {period_label(a.period_first)} / {period_label(a.period_last)}",
             _font(size=12, bold=True, color=GREEN), Alignment(wrap_text=False))
        notes = [f"{a.records_used} record attribuiti a un mese"]
        if a.records_spanning:
            notes.append(f"{a.records_spanning} con periodo plurimensile: valori ripartiti pro-quota sui mesi di competenza")
        if a.records_unplaced:
            notes.append(f"{a.records_unplaced} non attribuibili (periodo assente o superiore al limite) e quindi esclusi")
        _put(ws, row + 1, 1, " · ".join(notes), meta_font, Alignment(wrap_text=False))
        row += 3

        for m in a.matrices:
            years = list(m.years)
            unit = f" ({m.unit})" if m.unit else ""
            _put(ws, row, 1, f"{m.label}{unit}", _font(size=11, bold=True), Alignment(wrap_text=False))
            _header(ws, ["Mese"] + [str(y) for y in years], row=row + 1)
            top = row + 2
            r = top
            for mi in range(12):
                _put(ws, r, 1, MONTH_FULL[mi], data_font, Alignment(horizontal="left"), border)
                for c, y in enumerate(years, start=2):
                    v = m.total.values[y][mi]
                    _put(ws, r, c, v if v is not None else "-", data_font, Alignment(horizontal="center"), border)
                    if v is not None:
                        ws.cell(row=r, column=c).number_format = NUM_FMT
                r += 1
            summary = (("Totale", m.total.totals, NUM_FMT), ("Media mensile", m.total.monthly_avg, NUM_FMT),
                       ("Mesi con dato", m.total.months_with_data, "0"))
            for label, source, fmt in summary:
                _put(ws, r, 1, label, bold, Alignment(horizontal="left"), border, total_fill)
                for c, y in enumerate(years, start=2):
                    v = source.get(y)
                    _put(ws, r, c, v if v is not None else "-", bold, Alignment(horizontal="center"), border, total_fill)
                    if v is not None:
                        ws.cell(row=r, column=c).number_format = fmt
                r += 1
            if m.yoy:
                by_year = {y.year: y for y in m.yoy}
                for label, attr, fmt in (("Var. assoluta vs anno prec.", "delta", NUM_FMT),
                                         ("Var. % vs anno prec.", "delta_pct", '0.0"%"'),
                                         ("Mesi confrontati", "months_compared", "0")):
                    _put(ws, r, 1, label, data_font, Alignment(horizontal="left"), border)
                    for c, y in enumerate(years, start=2):
                        v = getattr(by_year[y], attr) if y in by_year else None
                        _put(ws, r, c, v if v is not None else "-", data_font, Alignment(horizontal="center"), border)
                        if v is not None:
                            ws.cell(row=r, column=c).number_format = fmt
                    r += 1
            if len(m.by_dimension) > 1:
                r += 1
                _header(ws, [a.dimension_label or "Dettaglio"] + [str(y) for y in years] + ["Totale"], row=r)
                r += 1
                for s_ in m.by_dimension:
                    _put(ws, r, 1, s_.name, data_font, Alignment(horizontal="left"), border)
                    for c, y in enumerate(years, start=2):
                        v = s_.totals.get(y)
                        _put(ws, r, c, v if v is not None else "-", data_font, Alignment(horizontal="center"), border)
                        if v is not None:
                            ws.cell(row=r, column=c).number_format = NUM_FMT
                    _put(ws, r, len(years) + 2, sum(s_.totals.values()), bold, Alignment(horizontal="center"), border)
                    ws.cell(row=r, column=len(years) + 2).number_format = NUM_FMT
                    r += 1
            max_cols = max(max_cols, len(years) + 2)

            chart_end = top
            png = render_monthly_chart(m, title=f"{a.label} — {m.label}{unit}")
            if png:
                img = XLImage(png)
                img.height = int(img.height * CHART_WIDTH_PX / img.width)
                img.width = CHART_WIDTH_PX
                ws.add_image(img, f"{get_column_letter(len(years) + 4)}{row}")
                chart_end = row + math.ceil(img.height / 20)
            row = max(r, chart_end) + 2
        row += 1

    ws.column_dimensions["A"].width = 30
    for c in range(2, max_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16


# ----------------------------------------------------------------------------- deadlines
def build_deadlines_from_payload(payload: Dict[str, Any], active_files: Optional[set] = None) -> List[Dict[str, Any]]:
    """
    Derive deadlines for ALL modules/schemas directly from the (reviewed) records.
    Used both for the report and for syncing the alerts DB, so they can never disagree.
    """
    out: List[Dict[str, Any]] = []
    seen = set()
    for mid, sname, rec in _iter_records(payload):
        src = rec.get("source_document")
        if active_files is not None and not _is_kept(src, active_files):
            continue
        try:
            items = extract_deadlines_from_records(mid, sname, [rec])
        except Exception:
            continue
        resp = _text((rec.get("fields") or {}).get("responsabile"))
        for d in items:
            if d["id"] in seen:
                continue
            seen.add(d["id"])
            d["responsabile"] = resp
            d["module_label"] = _mod_short(mid)
            out.append(d)
    return sorted(out, key=lambda x: x["due_date"])


def _next_notice(due: date, thresholds: Any, today: date) -> Optional[date]:
    if isinstance(thresholds, str):
        try:
            thresholds = json.loads(thresholds)
        except ValueError:
            thresholds = []
    cands = []
    for t in thresholds or []:
        try:
            nd = due - timedelta(days=int(t))
        except (TypeError, ValueError):
            continue
        if nd >= today:
            cands.append(nd)
    return min(cands) if cands else None


# ----------------------------------------------------------------------------- per-document metadata
def _doc_meta(entries: List[Tuple[str, str, Dict[str, Any]]]) -> Dict[str, Any]:
    subjects, issuers, resps, companies, addresses, issue_dates, due_dates = [], [], [], [], [], [], []
    cat_counts: Dict[str, int] = {}
    mods: List[str] = []
    n_edit = n_warn = 0

    for mid, sname, rec in entries:
        fields = rec.get("fields") or {}
        schema = _schema_for(mid, sname)
        label = schema.label if schema else sname
        cat_counts[label] = cat_counts.get(label, 0) + 1
        short = _mod_short(mid)
        if short not in mods:
            mods.append(short)

        for keys, bucket in ((SUBJECT_KEYS, subjects), (ISSUER_KEYS, issuers), (RESPONSIBLE_KEYS, resps),
                             (COMPANY_KEYS, companies), (ADDRESS_KEYS, addresses)):
            for k in keys:
                t = _text(fields.get(k))
                if t:
                    bucket.append(t)
                    break
        for k in ISSUE_DATE_KEYS:
            d = _date_of(fields.get(k))
            if d:
                issue_dates.append(d)
                break
        for n in _due_field_names(schema, rec):
            d = _date_of(fields.get(n))
            if d:
                due_dates.append(d)

        if any(isinstance(f, dict) and f.get("status") == "EDITED" for f in fields.values()):
            n_edit += 1
        if rec.get("warnings"):
            n_warn += 1

    return {
        "subject": _distinct(subjects),
        "issuer": _distinct(issuers),
        "responsabile": _distinct(resps),
        "company": _distinct(companies),
        "address": _distinct(addresses),
        "issue": min(issue_dates) if issue_dates else None,
        "due": min(due_dates) if due_dates else None,   # earliest = worst case for the document
        "cat_counts": cat_counts,
        "modules": mods,
        "n_records": len(entries),
        "n_edit": n_edit,
        "n_warn": n_warn,
    }


# ----------------------------------------------------------------------------- main
def create_platform_excel_report(
    documents: List[Dict[str, Any]],
    extracted_data_all_modules: Dict[str, Dict[str, Any]],
    deadlines: Optional[List[Dict[str, Any]]] = None,
    company_profile: Optional[Dict[str, Any]] = None,
    routing: Optional[Dict[str, str]] = None,
    errors: Optional[Dict[str, str]] = None,  # kept for API compatibility; the register no longer has a Note column
) -> bytes:
    today = date.today()
    comp = company_profile or {}
    profile_name = (comp.get("company_name") or "").strip()
    profile_address = (comp.get("address") or "").strip()
    comp_name = profile_name or "AZIENDA"
    comp_address = profile_address or "-"
    director = (comp.get("technical_director") or "").strip()

    names: List[str] = []
    for d in documents or []:
        n = d.get("filename") or d.get("name")
        if n and n not in names:
            names.append(n)
    active = set(names)

    payload = _filter_payload(extracted_data_all_modules, active)
    if deadlines is None:
        deadlines = build_deadlines_from_payload(payload)
    else:
        deadlines = [d for d in deadlines if _is_kept(d.get("source_document"), active)]

    by_doc: Dict[str, List[Tuple[str, str, Dict[str, Any]]]] = {}
    for mid, sname, rec in _iter_records(payload):
        src = rec.get("source_document")
        if src in active:
            by_doc.setdefault(src, []).append((mid, sname, rec))

    # ---- per-document rows
    doc_rows: List[Dict[str, Any]] = []
    for idx, name in enumerate(names, start=1):
        entries = by_doc.get(name, [])
        if entries:
            m = _doc_meta(entries)
            status, days = _calc_doc_status(m["due"], today)
            doc_rows.append({
                "idx": idx, "filename": name,
                # company / address: what the document itself says first, the tenant profile only as fallback
                "company": m["company"] if m["company"] != "-" else (profile_name or "-"),
                "address": m["address"] if m["address"] != "-" else (profile_address or "-"),
                "category": ", ".join(m["cat_counts"]), "project": ", ".join(m["modules"]),
                "subject": m["subject"], "issuer": m["issuer"],
                "issue": m["issue"], "due": m["due"], "status": status, "days": days,
                "resp": m["responsabile"] if m["responsabile"] != "-" else (director or "-"),
                "cat_counts": m["cat_counts"],
            })
        else:
            routed = (routing or {}).get(name)
            doc_rows.append({
                "idx": idx, "filename": name, "category": routed or "Non classificato", "project": "-",
                "company": profile_name or "-", "address": profile_address or "-",
                "subject": "-", "issuer": "-", "issue": None, "due": None,
                "status": "Non analizzato", "days": None, "resp": director or "-",
                "cat_counts": {},
            })

    count = lambda s: sum(1 for r in doc_rows if r["status"] == s)  # noqa: E731
    total, valid, expiring, expired = len(doc_rows), count("Valido"), count("In scadenza"), count("Scaduto")
    no_date, unanalyzed = count("Senza scadenza"), count("Non analizzato")
    dated = valid + expiring + expired
    compliance = f"{(valid + expiring) / dated * 100:.1f}%" if dated else "N/D"

    dl_days = [(_to_date(d.get("due_date")), d) for d in deadlines]
    dl_days = [((due - today).days, d) for due, d in dl_days if due]
    dl_overdue = sum(1 for dd, _ in dl_days if dd < 0)
    dl_30 = sum(1 for dd, _ in dl_days if 0 <= dd <= 30)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    border = _border()
    data_font = _font(size=FONT_DATA)
    meta_font = _font(size=FONT_DATA, italic=True, color="555555")

    # ======================================================= Sheet 1: Dashboard
    ws1 = wb.create_sheet("Dashboard & KPI")
    _put(ws1, 1, 1, f"DASHBOARD DOCUMENTALE & CONFORMITÀ — {comp_name.upper()}",
         _font(size=16, bold=True, color=GREEN))
    _put(ws1, 2, 1, f"Generato il: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Sede Operativa: {comp_address}",
         meta_font)

    kpi_fill = PatternFill(start_color="F4F6F4", end_color="F4F6F4", fill_type="solid")

    def kpi_block(row: int, items: List[Tuple[str, Any]]) -> None:
        for c, (label, val) in enumerate(items, start=1):
            _put(ws1, row, c, label, _font(size=FONT_KPI_LABEL, bold=True, color="555555"), Alignment(horizontal="center"))
            _put(ws1, row + 1, c, val, _font(size=FONT_KPI_VALUE, bold=True, color=GREEN),
                 Alignment(horizontal="center"), border, kpi_fill)
        ws1.row_dimensions[row].height = 24
        ws1.row_dimensions[row + 1].height = 38

    kpi_block(4, [("Totale Documenti", total), ("Documenti Validi", valid),
                  ("In Scadenza (≤ 30 gg)", expiring), ("Scaduti", expired),
                  ("% Conformità (non scaduti / con scadenza)", compliance)])
    kpi_block(7, [("Senza Scadenza", no_date), ("Non Analizzati", unanalyzed),
                  ("Scadenze Monitorate", len(deadlines)), ("Scadenze Scadute", dl_overdue),
                  ("Scadenze ≤ 30 gg", dl_30)])

    _put(ws1, 10, 1, "Dettaglio per Categoria:", _font(size=12, bold=True))
    _header(ws1, ["Categoria", "Documenti", "Record Estratti", "Scaduti", "In Scadenza"], row=11)
    cat_stats: Dict[str, Dict[str, int]] = {}
    for r in doc_rows:
        for cat, n in r["cat_counts"].items():
            s = cat_stats.setdefault(cat, {"docs": 0, "recs": 0, "exp": 0, "soon": 0})
            s["docs"] += 1
            s["recs"] += n
            s["exp"] += r["status"] == "Scaduto"
            s["soon"] += r["status"] == "In scadenza"
    row = 12
    for cat, s in cat_stats.items():
        for c, v in enumerate([cat, s["docs"], s["recs"], s["exp"], s["soon"]], start=1):
            _put(ws1, row, c, v, data_font, border=border)
        row += 1
    if not cat_stats:
        _put(ws1, 12, 1, "Nessuna estrazione eseguita: i documenti risultano 'Non analizzati'.", meta_font)

    # ======================================================= Sheet 2: Registro Documenti
    ws2 = wb.create_sheet("Registro Documenti")
    ws2.freeze_panes = "A2"
    reg_headers = [
        "ID Documento", "Azienda / Cliente", "Sede", "Progetto", "Categoria",
        "Titolo / Nome File", "Dipendente / Macchina / Asset", "Fornitore / Ente Emittente",
        "Data Emissione", "Data Scadenza", "Stato", "Responsabile",
    ]
    _header(ws2, reg_headers)
    for r_i, d in enumerate(doc_rows, start=2):
        values = [f"DOC-{d['idx']:03d}", d["company"], d["address"], d["project"], d["category"],
                  d["filename"], d["subject"], d["issuer"], d["issue"] or "-", d["due"] or "-",
                  d["status"], d["resp"]]
        for c, v in enumerate(values, start=1):
            _put(ws2, r_i, c, v, data_font, border=border)
        for c in (1, 9, 10, 11):
            ws2.cell(row=r_i, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.cell(row=r_i, column=11).font = _font(
            size=FONT_DATA, bold=d["status"] in STATUS_COLOR, color=STATUS_COLOR.get(d["status"], "555555"))
    _autofilter(ws2)

    # ======================================================= Sheet 3: Scadenze & Azioni
    ws3 = wb.create_sheet("Scadenze & Azioni")
    ws3.freeze_panes = "A2"
    _header(ws3, ["Documento Fonte", "Soggetto / Asset Collegato", "Data di Scadenza", "Giorni Mancanti",
                  "Data di Avviso", "Responsabile", "Priorità", "Stato", "Azione Richiesta"])
    r3 = 2
    for d in sorted(deadlines, key=lambda x: str(x.get("due_date", ""))):
        due = _to_date(d.get("due_date"))
        if due is None:
            continue
        status, days = _calc_doc_status(due, today)
        notice = _next_notice(due, d.get("thresholds"), today)
        prio = "Alta" if days <= 15 else ("Media" if days <= 30 else "Bassa")
        action = {"Scaduto": "Rinnovo / regolarizzazione urgente",
                  "In scadenza": "Pianificare rinnovo / aggiornamento"}.get(status, "Monitoraggio regolare")
        values = [d.get("source_document") or "-", d.get("description") or "-", due, days,
                  notice or "Soglie già raggiunte", d.get("responsabile") or director or "-",
                  prio, status, action]
        for c, v in enumerate(values, start=1):
            _put(ws3, r3, c, v, data_font, border=border)
        for c in (3, 4, 5, 7, 8):
            ws3.cell(row=r3, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws3.cell(row=r3, column=8).font = _font(size=FONT_DATA, bold=True, color=STATUS_COLOR.get(status, "2E7D32"))
        r3 += 1
    if r3 == 2:
        _put(ws3, 2, 1, "Nessuna scadenza rilevata nei record estratti.", meta_font)
    _autofilter(ws3)

    # ======================================================= Sheet 4+: Dati per schema
    calc_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    edit_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    unc_fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
    any_analytic = False

    for mid, schemas in payload.items():
        mod_obj = _module_for(mid)
        for sname, sdata in schemas.items():
            records = sdata["records"]
            any_analytic = True
            schema = _schema_for(mid, sname)
            ftypes = {n: f.type for n, f in schema.fields.items()} if schema else {}
            long_fields = {s["field"] for s in (schema.word_sections if schema else [])}

            fallback = f"Dati {sname.replace('_', ' ').title()}"
            ws = wb.create_sheet(_unique_title(wb, sheet_title_for(schema, sname, fallback)))
            ws.freeze_panes = "A2"
            plan = column_plan(sname, mod_obj, records)
            expand = schema.excel_expand if schema else {}
            keyed = expand.get("keyed_fields") or {}
            with_status = not (schema and schema.hide_status_column)

            # Column layout: one header per plan field; per-code fields that declare a unit get a unit column right after.
            headers = ["Documento Fonte", "Pagina Fonte"]
            for fname, flabel in plan:
                headers.append(flabel)
                if fname in keyed and keyed[fname].get("unit_label"):
                    headers.append(keyed[fname]["unit_label"])
            if with_status:
                headers.append("Stato")
            _header(ws, headers)

            r_i = 2
            for rec in records:
                fields = rec.get("fields") or {}
                # one report row per code when the schema asks for it, else exactly one row for the record
                variants = expand_record({n: _text(fields.get(n)) for n in [expand["code_field"], *keyed]}, expand)                     if expand else [None]
                for variant in variants:
                    c = 1
                    _put(ws, r_i, c, rec.get("source_document") or "-", data_font, border=border)
                    c += 1
                    _put(ws, r_i, c, f"Pagina {rec.get('source_page') or '-'}",
                         _font(size=FONT_DATA, italic=True, color="555555"), Alignment(horizontal="center"), border)
                    for fname, _ in plan:
                        c += 1
                        fd = fields.get(fname) or {}
                        st = fd.get("status")
                        fill = calc_fill if st == "CALCULATED" else edit_fill if st == "EDITED" else                             unc_fill if st == "UNCERTAIN" else None
                        if variant is not None and (fname == expand["code_field"] or fname in keyed):
                            val = variant.get(fname)
                            # codes, quantities and operations are short tokens: centred
                            _put(ws, r_i, c, val if val not in (None, "") else "-", data_font,
                                 Alignment(horizontal="center"), border, fill)
                            if fname in keyed and keyed[fname].get("unit_label"):
                                c += 1
                                unit = variant.get(unit_key(fname))
                                _put(ws, r_i, c, unit or "-", data_font, Alignment(horizontal="center"), border, fill)
                            continue
                        val = cell_value(fd)
                        if ftypes.get(fname) == "date" and _to_date(val):
                            val = _to_date(val)
                        align = Alignment(wrap_text=True, vertical="top") if fname in long_fields else None
                        _put(ws, r_i, c, val if val is not None else "-", data_font, align=align, border=border, fill=fill)
                    if with_status:
                        _put(ws, r_i, c + 1, _record_status(rec), data_font, Alignment(horizontal="center"), border)
                    r_i += 1
            _autofilter(ws)
            _autosize(ws)
            _fit_row_heights(ws)

    if not any_analytic:
        ws_e = wb.create_sheet("Dati Estratti (Analitici)")
        _put(ws_e, 1, 1, "Nessun dato analitico estratto. Eseguire l'estrazione nei rispettivi moduli.", meta_font)

    analyses = analyze_payload(payload)          # [] for documents without consumption data -> no sheet, no chart
    if analyses:
        write_trend_sheet(wb, analyses)

    for ws_item in (ws2, ws3):
        _autosize(ws_item)
        _fit_row_heights(ws_item)
    for r in range(12, ws1.max_row + 1):
        ws1.row_dimensions[r].height = ROW_HEIGHT
    for letter, w in zip("ABCDE", (34, 22, 22, 22, 22)):
        ws1.column_dimensions[letter].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()