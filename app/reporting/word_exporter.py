import io
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.reporting.analytics import ConsumptionAnalysis, analyze_payload, format_it, period_label
from app.reporting.charts import render_monthly_chart
from app.reporting.common import (
    resolve_scope,
    report_meta,
    column_plan,
    cell_value,
    signoff_labels,
    split_items,
)


def _set_cell_shading(cell, color_hex: str):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def _style_table_grid(table, border_color: str = "C8D6C5"):
    tbl_pr = table._tbl.tblPr
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'</w:tblBorders>'
    )
    cell_mar_xml = (
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="90" w:type="dxa"/>'
        f'<w:left w:w="130" w:type="dxa"/>'
        f'<w:bottom w:w="90" w:type="dxa"/>'
        f'<w:right w:w="130" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tbl_pr.append(parse_xml(borders_xml))
    tbl_pr.append(parse_xml(cell_mar_xml))


TREND_SECTION_TITLE = "Analisi Grafica dei Consumi e Trend Economico"


def _plain_cell(cell, text: str, bold: bool = False, align=None, size: float = 9, color: Optional[RGBColor] = None):
    cell.text = ""
    par = cell.paragraphs[0]
    if align is not None:
        par.alignment = align
    run = par.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _summary_rows(a: ConsumptionAnalysis) -> List[Tuple[str, str]]:
    """Management summary: period, total cost, total volume and peak month — all straight from the aggregated matrices."""
    rows = [("Periodo analizzato", f"{period_label(a.period_first)} – {period_label(a.period_last)}")]
    for m in a.matrices:
        unit = f" {m.unit}" if m.unit else ""
        if m.kind == "cost":
            rows.append((f"Spesa totale del periodo — {m.label}", f"{format_it(m.grand_total)}{unit}"))
        else:
            rows.append((f"Consumo totale — {m.label}", f"{format_it(m.grand_total)}{unit}"))
            if m.peak:
                y, mo, v = m.peak
                rows.append((f"Mese di picco — {m.label}", f"{period_label((y, mo))} ({format_it(v)}{unit})"))
    rows.append(("Record considerati", str(a.records_used)))
    if a.records_spanning:
        rows.append(("Periodi plurimensili", f"{a.records_spanning} (valori ripartiti pro-quota sui mesi di competenza)"))
    if a.records_unplaced:
        rows.append(("Record esclusi", f"{a.records_unplaced} (periodo assente o non attribuibile a un singolo mese)"))
    return rows


def _add_trend_section(doc, analyses: List[ConsumptionAnalysis], sec_no: int, color: RGBColor) -> None:
    h = doc.add_heading(level=1)
    run_h = h.add_run(f"{sec_no}. {TREND_SECTION_TITLE}")
    run_h.font.name = "Segoe UI"
    run_h.font.color.rgb = color

    for a in analyses:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(4)
        p_sub.paragraph_format.space_after = Pt(4)
        r_sub = p_sub.add_run(a.label)
        r_sub.font.name = "Segoe UI"
        r_sub.font.size = Pt(11)
        r_sub.font.bold = True
        r_sub.font.color.rgb = color

        for m in a.matrices:
            unit = f" ({m.unit})" if m.unit else ""
            png = render_monthly_chart(m, title=f"{m.label}{unit}")
            if png is None:
                continue
            frame = doc.add_table(rows=1, cols=1)        # corporate frame around the picture
            frame.alignment = WD_TABLE_ALIGNMENT.CENTER
            _style_table_grid(frame, "8CC63F")
            cell = frame.rows[0].cells[0]
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(png, width=Inches(6.3))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _style_table_grid(table, "C8D6C5")
        for cell, text in zip(table.rows[0].cells, ("Sintesi gestionale", "Valore")):
            _set_cell_shading(cell, "1E4D2B")
            _plain_cell(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, color=RGBColor(255, 255, 255))
        for label, value in _summary_rows(a):
            cells = table.add_row().cells
            _set_cell_shading(cells[0], "EFF5EE")
            _plain_cell(cells[0], label, bold=True, size=9, color=color)
            _plain_cell(cells[1], value, size=9)
        doc.add_paragraph().paragraph_format.space_after = Pt(8)


def create_styled_word_report(
    extracted_data_by_schema: Dict[str, Any],
    source_documents: List[str],
    company_profile: Optional[Dict[str, Any]] = None,
    selected_fields_by_schema: Optional[Dict[str, List[str]]] = None,
    module_id: Optional[str] = None
) -> bytes:
    """
    تولید سند رسمی و مدیریتی Word (.docx) با تفکیک قطعی ماژول، کادرهای شناسنامه‌ای عمودی (Key-Value) و جداول بهینه.
    """
    module, scoped_data = resolve_scope(extracted_data_by_schema, module_id)
    meta = report_meta(module)

    doc = Document()
    comp = company_profile or {}
    comp_name = comp.get("company_name", "AZIENDA")
    vat_number = comp.get("vat_number", "")
    address = comp.get("address", "")

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    COLOR_PRIMARY = RGBColor(30, 77, 43)
    COLOR_SECONDARY = RGBColor(85, 85, 85)

    # تعیین عنوان پویای هدر گزارش بر اساس اسکیمای فعال (schema_obj.report_title)
    active_schemas = [
        module.schemas[s] for s, d in scoped_data.items()
        if module and s in module.schemas and d.get("records")
    ] or ([module.schemas[s] for s in scoped_data if module and s in module.schemas] if module else [])
    schema_obj = next((s for s in active_schemas if s.layout.get("report_title")), active_schemas[0] if active_schemas else None)
    if schema_obj and schema_obj.layout.get("report_title"):
        doc_title_text = schema_obj.report_title
    elif meta.get("doc_title") and meta.get("doc_title") != "RAPPORTO DI CONFORMITÀ E AUDIT":
        doc_title_text = meta["doc_title"]
    elif schema_obj:
        doc_title_text = schema_obj.report_title
    else:
        doc_title_text = meta.get("doc_title", "RAPPORTO DI CONFORMITÀ E AUDIT")

    # عنوان اصلی
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run(f"{doc_title_text} — {comp_name.upper()}")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # زیرعنوان
    sub_title = meta.get("doc_subtitle", "")
    if sub_title:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(4)
        run_sub = sub_p.add_run(sub_title)
        run_sub.font.name = "Segoe UI"
        run_sub.font.size = Pt(10.5)
        run_sub.font.bold = True
        run_sub.font.color.rgb = COLOR_SECONDARY

    if vat_number or address:
        p_cinfo = doc.add_paragraph()
        p_cinfo.paragraph_format.space_after = Pt(10)
        c_line = f"P.IVA / C.F.: {vat_number}" if vat_number else ""
        if address:
            c_line += f"  |  Sede Operativa: {address}"
        run_cinfo = p_cinfo.add_run(c_line)
        run_cinfo.font.name = "Segoe UI"
        run_cinfo.font.size = Pt(9.5)
        run_cinfo.font.italic = True
        run_cinfo.font.color.rgb = COLOR_SECONDARY

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    run_div = p_div.add_run("―" * 58)
    run_div.font.color.rgb = RGBColor(200, 200, 200)

    # بخش ۱: متادیتای ممیزی و اسناد
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Informazioni Generali e Documenti Analizzati")
    run_h1.font.name = "Segoe UI"
    run_h1.font.color.rgb = COLOR_PRIMARY

    p_meta = doc.add_paragraph()
    p_meta.add_run("Data di Generazione: ").bold = True
    p_meta.add_run(f"{datetime.now().strftime('%d/%m/%Y ore %H:%M:%S')}\n")
    p_meta.add_run("Metodologia: ").bold = True
    p_meta.add_run("Analisi multimodale semantica con verifica anti-allucinazione e tracciabilità di pagina.\n")
    p_meta.add_run("Documenti di Origine Verificati:\n").bold = True
    
    for doc_name in source_documents:
        p_doc = doc.add_paragraph(style="List Bullet")
        p_doc.paragraph_format.space_after = Pt(2)
        r = p_doc.add_run(doc_name)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # بخش ۲: مشخصات و جداول استخراج‌شده به تفکیک اسکیما
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. Tabelle di Dettaglio per Schema")
    run_h2.font.name = "Segoe UI"
    run_h2.font.color.rgb = COLOR_PRIMARY

    for schema_name, data in scoped_data.items():
        records = data.get("records", [])
        schema_obj = module.schemas.get(schema_name) if module else None
        schema_label = schema_obj.label if schema_obj else schema_name.upper()

        p_schema_title = doc.add_paragraph()
        p_schema_title.paragraph_format.space_before = Pt(8)
        p_schema_title.paragraph_format.space_after = Pt(4)
        run_sname = p_schema_title.add_run(f"Sezione: {schema_label} — {schema_name.upper()} (Totale Record: {len(records)})")
        run_sname.font.name = "Segoe UI"
        run_sname.font.size = Pt(11.5)
        run_sname.font.bold = True

        if not records:
            p_empty = doc.add_paragraph("Nessun record identificato nei documenti per questo schema.")
            p_empty.runs[0].font.italic = True
            continue

        user_selected = (selected_fields_by_schema or {}).get(schema_name)
        full_col_plan = column_plan(schema_name, module, records, user_selected)
        section_fields = {s["field"] for s in (schema_obj.word_sections if schema_obj else [])}
        header_field_names = schema_obj.word_header_fields if schema_obj else []
        table_field_names = schema_obj.word_table_fields if schema_obj else []

        # الف) رندر بلوک شناسنامه‌ای عمودی (Key-Value) در صورت تعریف word_header_fields
        if header_field_names:
            all_labels_map = schema_obj.all_labels() if schema_obj else {}
            header_cols = [
                (fn, all_labels_map.get(fn, fn))
                for fn in header_field_names
                if not user_selected or fn in user_selected
            ]
            for r_idx, rec in enumerate(records, start=1):
                fields = rec.get("fields", {})
                if len(records) > 1:
                    p_rec_hdr = doc.add_paragraph()
                    p_rec_hdr.paragraph_format.space_before = Pt(4)
                    p_rec_hdr.paragraph_format.space_after = Pt(2)
                    r_rh = p_rec_hdr.add_run(f"Scheda Identificativa Record #{r_idx} (Pag. {rec.get('source_page', 1)})")
                    r_rh.font.name = "Segoe UI"
                    r_rh.font.size = Pt(9.5)
                    r_rh.font.bold = True
                    r_rh.font.color.rgb = COLOR_PRIMARY

                kv_table = doc.add_table(rows=0, cols=2)
                kv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                kv_table.autofit = False
                _style_table_grid(kv_table, "C8D6C5")

                for fname, flabel in header_cols:
                    val_disp = cell_value(fields.get(fname, {}))
                    row_cells = kv_table.add_row().cells
                    row_cells[0].width = Inches(2.3)
                    row_cells[1].width = Inches(4.5)
                    _set_cell_shading(row_cells[0], "EFF5EE")

                    p_key = row_cells[0].paragraphs[0]
                    r_key = p_key.add_run(flabel)
                    r_key.font.name = "Segoe UI"
                    r_key.font.bold = True
                    r_key.font.size = Pt(9)
                    r_key.font.color.rgb = COLOR_PRIMARY

                    p_val = row_cells[1].paragraphs[0]
                    r_val = p_val.add_run(str(val_disp if val_disp is not None else "-"))
                    r_val.font.name = "Segoe UI"
                    r_val.font.size = Pt(9)

                    row_cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    row_cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # ب) جدول افقی داده‌های کلیدی / پسماندها و مقادیر (بدون ستون Stato و بدون فیلدهای متنی طولانی)
        if table_field_names:
            table_col_plan = [c for c in full_col_plan if c[0] in set(table_field_names)]
            include_source_page = False
        else:
            excluded = section_fields | set(header_field_names)
            table_col_plan = [c for c in full_col_plan if c[0] not in excluded]
            include_source_page = True

        if table_col_plan:
            headers = (["Pagina Fonte"] if include_source_page else []) + [col[1] for col in table_col_plan]

            table = doc.add_table(rows=1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            _style_table_grid(table, "C8D6C5")

            hdr_cells = table.rows[0].cells
            for idx, text in enumerate(headers):
                hdr_cells[idx].text = text
                _set_cell_shading(hdr_cells[idx], "1E4D2B")
                hdr_p = hdr_cells[idx].paragraphs[0]
                hdr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in hdr_p.runs:
                    run.font.name = "Segoe UI"
                    run.font.bold = True
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)

            for rec in records:
                row_cells = table.add_row().cells
                src_page = rec.get("source_page", 1)
                fields = rec.get("fields", {})

                col_offset = 0
                if include_source_page:
                    row_cells[0].text = f"Pag. {src_page}"
                    row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    col_offset = 1

                for c_idx, (fname, flabel) in enumerate(table_col_plan, start=col_offset):
                    fdata = fields.get(fname, {})
                    val_to_display = cell_value(fdata)
                    row_cells[c_idx].text = str(val_to_display if val_to_display is not None else "-")

                for cell in row_cells:
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = "Segoe UI"
                            run.font.size = Pt(8.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # بخش‌های اختصاصی Word تعریف‌شده در layout.word_sections اسکیماها
    sec_no = 2
    for schema_name, data in scoped_data.items():
        schema_obj = module.schemas.get(schema_name) if module else None
        if not schema_obj or not schema_obj.word_sections:
            continue
        records = data.get("records", [])
        for sec in schema_obj.word_sections:
            sec_no += 1
            h_sec = doc.add_heading(level=1)
            run_hsec = h_sec.add_run(f"{sec_no}. {sec['title']}")
            run_hsec.font.name = "Segoe UI"
            run_hsec.font.color.rgb = COLOR_PRIMARY

            items: List[Tuple[str, str]] = []
            seen_items = set()
            for r in records:
                src_doc = r.get("source_document") or "-"
                fdict = (r.get("fields") or {}).get(sec["field"]) or {}
                raw_txt = fdict.get("normalized_value") or fdict.get("raw_value")
                for it in split_items(raw_txt):
                    if it not in seen_items:
                        seen_items.add(it)
                        items.append((it, src_doc))

            if not items:
                p_none = doc.add_paragraph("Nessuna prescrizione rilevata.")
                p_none.runs[0].font.italic = True
                p_none.runs[0].font.name = "Segoe UI"
                p_none.runs[0].font.size = Pt(9.5)
            elif sec["style"] == "table":
                sec_headers = ["N.", "Prescrizione / Condizione", "Scadenza / Termine", "Documento Fonte"]
                sec_table = doc.add_table(rows=1, cols=4)
                sec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                sec_table.autofit = True
                _style_table_grid(sec_table, "C8D6C5")

                shdr_cells = sec_table.rows[0].cells
                for idx, text in enumerate(sec_headers):
                    shdr_cells[idx].text = text
                    _set_cell_shading(shdr_cells[idx], "1E4D2B")
                    hp = shdr_cells[idx].paragraphs[0]
                    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in hp.runs:
                        run.font.name = "Segoe UI"
                        run.font.bold = True
                        run.font.size = Pt(8.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)

                for n, (it, src_doc) in enumerate(items, start=1):
                    cond, sep_found, due = it.partition("||")
                    cond_txt = cond.strip() or "-"
                    due_txt = due.strip() if sep_found and due.strip() else "-"
                    rcells = sec_table.add_row().cells
                    rcells[0].text = str(n)
                    rcells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rcells[1].text = cond_txt
                    rcells[2].text = due_txt
                    rcells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rcells[3].text = src_doc
                    for cell in rcells:
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.name = "Segoe UI"
                                run.font.size = Pt(8.5)
            elif sec["style"] == "list":
                for n, (it, src_doc) in enumerate(items, start=1):
                    p_item = doc.add_paragraph()
                    p_item.paragraph_format.space_after = Pt(3)
                    suffix = f" ({src_doc})" if src_doc and src_doc != "-" else ""
                    r_item = p_item.add_run(f"{n}. {it}{suffix}")
                    r_item.font.name = "Segoe UI"
                    r_item.font.size = Pt(9.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # تحلیل گرافیکی مصرف/هزینه: فقط اگر رکوردهای مصرفی معتبر وجود داشته باشند (در غیر این صورت هیچ بخشی درج نمی‌شود)
    analyses = analyze_payload({module.module_id: scoped_data}) if module else []
    if analyses:
        sec_no += 1
        _add_trend_section(doc, analyses, sec_no, COLOR_PRIMARY)

    # بخش پایانی: کادرهای امضا (داینامیک بر اساس ماژول)
    signs = signoff_labels(meta, comp)
    if signs:
        h_sign = doc.add_heading(level=1)
        run_hsign = h_sign.add_run(f"{sec_no + 1}. Approvazione e Visto")
        run_hsign.font.name = "Segoe UI"
        run_hsign.font.color.rgb = COLOR_PRIMARY

        sign_table = doc.add_table(rows=2, cols=len(signs))
        sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _style_table_grid(sign_table, "C8D6C5")
        
        for idx, sign_txt in enumerate(signs):
            sign_table.rows[0].cells[idx].text = sign_txt
            sign_table.rows[1].cells[idx].text = "Firma: _________________________\nData:   ____ / ____ / ________"

        for row in sign_table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Segoe UI"
                        run.font.size = Pt(9.5)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()

