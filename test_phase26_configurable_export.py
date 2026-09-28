import io
import openpyxl
from docx import Document
from app.extraction.schemas import AVAILABLE_SCHEMAS
from app.reporting.excel_exporter import create_styled_excel_report
from app.reporting.word_exporter import create_styled_word_report


def test_phase26_configurable_export():
    print("Testing Phase 26 Configurable Export Engine & Updated Schemas...")

    # ۱. اعتبارسنجی فیلدهای جدید اسکیمای آموزش
    p_fields = AVAILABLE_SCHEMAS["personnel_training"]["fields"]
    assert "codice_fiscale" in p_fields
    assert "riferimento_normativo" in p_fields
    assert "numero_protocollo" in p_fields
    print("✓ Personnel training schema verified with Codice Fiscale, Protocol, and Legal Reference.")

    # ۲. داده ساختگی با تمام فیلدها
    test_data = {
        "personnel_training": {
            "total_records": 1,
            "records": [
                {
                    "source_page": 1,
                    "fields": {
                        "nominativo_dipendente": {"normalized_value": "Ermes Frossasco", "status": "EXTRACTED"},
                        "codice_fiscale": {"normalized_value": "FRSRMS90M03I470K", "status": "EXTRACTED"},
                        "corso_descrizione": {"normalized_value": "Uso Attrezzature Lavoro", "status": "EXTRACTED"},
                        "riferimento_normativo": {"normalized_value": "Artt. 71 e 73 D.Lgs. 81/08", "status": "EXTRACTED"},
                        "numero_protocollo": {"normalized_value": "EB00F763/2026/0655", "status": "EXTRACTED"},
                        "ente_formatore": {"normalized_value": "Conflavoro PMI / 3R International", "status": "EXTRACTED"},
                        "data_emissione": {"normalized_value": "2026-04-22", "status": "EXTRACTED"},
                        "data_scadenza_rinnovo": {"normalized_value": "2031-04-22", "status": "EXTRACTED"},
                        "frequenza_anni": {"normalized_value": 5.0, "status": "EXTRACTED"},
                        "ore_formazione": {"normalized_value": 4.0, "status": "EXTRACTED"},
                        "stato_validita": {"normalized_value": "Valido", "status": "EXTRACTED"}
                    }
                }
            ]
        }
    }

    # ۳. انتخاب ۳ ستون خاص توسط کاربر (حذف سایر ستون‌ها در خروجی)
    custom_selection = {
        "personnel_training": ["nominativo_dipendente", "codice_fiscale", "data_scadenza_rinnovo"]
    }

    # تولید اکسل با ستون‌های فیلترشده
    excel_bytes = create_styled_excel_report(
        test_data,
        ["attestato.pdf"],
        selected_fields_by_schema=custom_selection
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb["PERSONNEL_TRAINING"]
    headers = [cell.value for cell in ws[1]]
    assert "Dipendente / Collaboratore" in headers
    assert "Codice Fiscale" in headers
    assert "Scadenza Aggiornamento" in headers
    assert "Ente Formatore / CFPT" not in headers  # ستونی که کاربر حذف کرده نباید وجود داشته باشد
    print("✓ Excel export dynamically filtered columns according to user selection.")

    # تولید ورد با ستون‌های فیلترشده
    word_bytes = create_styled_word_report(
        test_data,
        ["attestato.pdf"],
        selected_fields_by_schema=custom_selection
    )
    doc = Document(io.BytesIO(word_bytes))
    table = doc.tables[0]
    word_headers = [c.text for c in table.rows[0].cells]
    assert "Dipendente / Collaboratore" in word_headers
    assert "Codice Fiscale" in word_headers
    assert "Ente Formatore / CFPT" not in word_headers
    print("✓ Word report table dynamically filtered columns according to user selection.")


if __name__ == "__main__":
    test_phase26_configurable_export()
    print("\n🎉 PHASE 26 CONFIGURABLE EXPORT TESTS PASSED!")