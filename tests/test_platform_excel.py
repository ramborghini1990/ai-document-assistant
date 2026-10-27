import io
import openpyxl
import pytest
from app.reporting.platform_excel_exporter import create_platform_excel_report


@pytest.mark.offline
def test_platform_excel_4_sheets_structure():
    sample_docs = [{"filename": "bolletta_test.pdf"}]
    sample_extracted_all = {
        "energy_utility": {
            "energy_bill": {
                "records": [{
                    "source_document": "bolletta_test.pdf",
                    "source_page": 1,
                    "fields": {
                        "codice_pod": {"raw_value": "IT001E123", "normalized_value": "IT001E123", "status": "EXTRACTED"},
                        "consumo_kwh": {"raw_value": "1000", "normalized_value": 1000.0, "status": "EXTRACTED"},
                        "data_scadenza": {"raw_value": "2026-11-30", "normalized_value": "2026-11-30", "status": "EXTRACTED"}
                    }
                }]
            }
        }
    }
    company_prof = {"company_name": "Studio & Impresa S.r.l.", "address": "Torino"}

    excel_bytes = create_platform_excel_report(
        documents=sample_docs,
        extracted_data_all_modules=sample_extracted_all,
        company_profile=company_prof
    )

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Dashboard & KPI" in wb.sheetnames
    assert "Registro Documenti" in wb.sheetnames
    assert "Scadenze & Azioni" in wb.sheetnames
    assert "Dati Energy Bill" in wb.sheetnames
    # schemas without `hide_status_column` keep the record-status column
    assert [c.value for c in wb["Dati Energy Bill"][1]][-1] == "Stato"


def _f(value, status="EXTRACTED"):
    return {"raw_value": value, "normalized_value": value, "status": status}


def _authorization_workbook(fields, profile=None):
    payload = {"iso_14001": {"regulatory_authorization": {"records": [{
        "source_document": "104-2531_2023.pdf", "source_page": 1, "fields": fields}]}}}
    out = create_platform_excel_report(
        documents=[{"filename": "104-2531_2023.pdf"}],
        extracted_data_all_modules=payload,
        company_profile=profile if profile is not None else {"company_name": "inambiente S.r.l.", "address": "Torino"},
    )
    return openpyxl.load_workbook(io.BytesIO(out))


@pytest.mark.offline
def test_platform_excel_regulatory_authorization_expands_one_row_per_eer_code():
    wb = _authorization_workbook({
        "numero_atto_protocollo": _f("104-2531"),
        "ente_rilascio": _f("Provincia"),
        "elenco_codici_eer": _f("16 01 03, 16.01.04*"),
        "quantitativi_movimentabili": _f("16.01.03: 9 t/anno\n16.01.04*: 1 t/anno"),
        "quantitativi_deposito": _f("16.01.03: 4 t; 16.01.04*: 5 t"),
        "operazioni_trattamento": _f("16.01.03: R13\n16.01.04*: R13, R12"),
        "prescrizioni_con_scadenza": _f("Invio piano || entro 90 giorni"),
    })
    assert "Dettagli File" in wb.sheetnames
    assert "Dati Regulatory Authorization" not in wb.sheetnames

    ws = wb["Dettagli File"]
    headers = [c.value for c in ws[1]]
    i = headers.index("Codici EER / Rifiuti Ammessi")
    assert headers[i + 1:i + 6] == [
        "Quantitativi autorizzati movimentabili annui", "U. di misura",
        "Quantitativi autorizzati in deposito", "U. di misura", "Operazioni di trattamento",
    ]
    assert "Stato" not in headers

    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 2                                   # one row per EER code
    assert rows[0][i:i + 6] == ["16.01.03", 9, "t/anno", 4, "t", "R13"]
    assert rows[1][i:i + 6] == ["16.01.04*", 1, "t/anno", 5, "t", "R13, R12"]
    j = headers.index("N. Atto / Protocollo")
    assert rows[0][j] == rows[1][j] == "104-2531"           # document-level columns repeated on every row
    assert rows[0][0] == rows[1][0] == "104-2531_2023.pdf"

    col_psc = headers.index("Prescrizioni con scadenza") + 1
    assert ws.cell(row=2, column=col_psc).alignment.wrap_text is True
    assert ws.cell(row=2, column=col_psc).alignment.vertical == "top"


@pytest.mark.offline
def test_platform_excel_global_figures_without_codes_do_not_leak_onto_every_row():
    wb = _authorization_workbook({
        "elenco_codici_eer": _f("15 01 01, 17 04 05"),
        "quantitativi_movimentabili": _f("Fase 1: 3000 t/a"),
        "quantitativi_deposito": _f("Fase 1: 150 t"),
    })
    ws = wb["Dettagli File"]
    headers = [c.value for c in ws[1]]
    i = headers.index("Codici EER / Rifiuti Ammessi")
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[i] for r in rows] == ["15.01.01", "17.04.05"]
    assert all(r[i + 1] == "-" and r[i + 3] == "-" for r in rows)


@pytest.mark.offline
def test_platform_excel_registry_uses_extracted_company_and_address_and_has_no_note_column():
    wb = _authorization_workbook({
        "ragione_sociale": _f("ALPHA S.R.L."),
        "sede_indirizzo": _f("VIA ROMA, 100 - BORGARO (TO)"),
        "numero_atto_protocollo": _f("104-2531"),
    }, profile={"company_name": "Profilo S.p.A.", "address": "Altrove"})
    ws = wb["Registro Documenti"]
    headers = [c.value for c in ws[1]]
    assert "Note" not in headers and headers[-1] == "Responsabile"
    row = {h: ws.cell(row=2, column=n).value for n, h in enumerate(headers, start=1)}
    assert row["Azienda / Cliente"] == "ALPHA S.R.L."
    assert row["Sede"] == "VIA ROMA, 100 - BORGARO (TO)"


@pytest.mark.offline
def test_platform_excel_registry_falls_back_to_profile_then_dash():
    wb = _authorization_workbook({"numero_atto_protocollo": _f("104-2531")},
                                 profile={"company_name": "Profilo S.p.A.", "address": ""})
    ws = wb["Registro Documenti"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("Azienda / Cliente") + 1).value == "Profilo S.p.A."
    assert ws.cell(row=2, column=headers.index("Sede") + 1).value == "-"
