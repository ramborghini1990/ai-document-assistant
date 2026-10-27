"""Offline checks for the consumption analytics, the charts and their Excel / Word integration."""
import io
from datetime import date

import openpyxl
import pytest
from docx import Document
from PIL import Image

from app.extraction.registry import _parse_schema, resolve_schema
from app.reporting.analytics import (
    NO_DIMENSION, analyze_payload, analyze_records, format_it, parse_period,
)
from app.reporting.charts import render_monthly_chart
from app.reporting.platform_excel_exporter import TREND_SHEET, create_platform_excel_report
from app.reporting.word_exporter import TREND_SECTION_TITLE, create_styled_word_report


def _f(value, status="EXTRACTED"):
    return {"raw_value": value, "normalized_value": value, "status": status}


def bill(period, kwh, eur=None, pod="IT001", kwh_status="EXTRACTED"):
    fields = {"periodo_riferimento": _f(period), "consumo_kwh": _f(kwh, kwh_status),
              "codice_pod": _f(pod)}
    if eur is not None:
        fields["costo_totale_eur"] = _f(eur)
    return {"source_document": "bolletta.pdf", "source_page": 1, "fields": fields}


def energy_schema():
    return resolve_schema("energy_bill", "energy_utility")


def matrix_of(analysis, field):
    return next(m for m in analysis.matrices if m.field == field)


# ----------------------------------------------------------------------------- period parsing
@pytest.mark.offline
@pytest.mark.parametrize("value,expected", [
    ("Settembre 2025", ((2025, 9), (2025, 9))),
    ("set 2025", ((2025, 9), (2025, 9))),
    ("15 settembre 2025", ((2025, 9), (2025, 9))),
    ("31/08/2025 - 30/09/2025", ((2025, 8), (2025, 9))),
    ("dal 01.10.2025 al 31.10.2025", ((2025, 10), (2025, 10))),
    ("2025-09-30", ((2025, 9), (2025, 9))),
    (date(2025, 3, 5), ((2025, 3), (2025, 3))),
    ("09/2025", ((2025, 9), (2025, 9))),
    ("2025-09", ((2025, 9), (2025, 9))),
    ("Luglio - Agosto 2025", ((2025, 7), (2025, 8))),
    ("Nov - Gen 2025", ((2024, 11), (2025, 1))),
    ("Dicembre 2024 - Gennaio 2025", ((2024, 12), (2025, 1))),
    ("Settembre", None), ("2025", None), ("Q3 2025", None), ("31/13/2025", None), ("", None), (None, None), (42, None),
])
def test_parse_period_understands_real_formats_and_never_guesses(value, expected):
    assert parse_period(value) == expected


@pytest.mark.offline
def test_italian_number_format():
    assert format_it(1234.5) == "1.234,50" and format_it(0, 0) == "0" and format_it(None) == "-"


# ----------------------------------------------------------------------------- aggregation
@pytest.mark.offline
def test_monthly_matrix_totals_averages_peak_and_like_for_like_yoy():
    recs = [bill("Settembre 2024", 800, 200), bill("Ottobre 2024", 1000, 300), bill("Dicembre 2024", 500, 100),
            bill("Settembre 2025", 877, 230.5), bill("Ottobre 2025", 950, 260), bill("Gennaio 2025", 1500, 520)]
    a = analyze_records("energy_utility", energy_schema(), recs)
    kwh = matrix_of(a, "consumo_kwh")
    assert kwh.unit == "kWh" and kwh.years == (2024, 2025) and kwh.kind == "quantity"
    assert kwh.total.values[2024][8] == 800 and kwh.total.values[2024][0] is None      # missing month stays None
    assert kwh.total.totals == {2024: 2300, 2025: 3327} and kwh.grand_total == 5627
    assert kwh.total.monthly_avg[2024] == pytest.approx(2300 / 3)
    assert kwh.total.months_with_data == {2024: 3, 2025: 3}
    assert kwh.peak == (2025, 1, 1500)
    (yoy,) = kwh.yoy                                      # only Sep and Oct exist in both years
    assert (yoy.year, yoy.base_year, yoy.months_compared) == (2025, 2024, 2)
    assert yoy.value == 1827 and yoy.base_value == 1800 and yoy.delta == 27 and yoy.delta_pct == pytest.approx(1.5)
    cost = matrix_of(a, "costo_totale_eur")
    assert cost.kind == "cost" and cost.unit == "€" and cost.grand_total == pytest.approx(1610.5)
    assert (a.period_first, a.period_last) == ((2024, 9), (2025, 10)) and a.records_used == 6


@pytest.mark.offline
def test_dimensions_are_summed_into_the_total_and_missing_dimension_is_named():
    recs = [bill("Settembre 2025", 100, pod="IT002"), bill("Settembre 2025", 50, pod="IT001"),
            bill("Settembre 2025", 25, pod="")]
    a = analyze_records("energy_utility", energy_schema(), recs)
    kwh = matrix_of(a, "consumo_kwh")
    assert [s.name for s in kwh.by_dimension] == ["IT001", "IT002", NO_DIMENSION]
    assert kwh.total.values[2025][8] == 175
    assert a.dimension_label == "POD"


@pytest.mark.offline
def test_only_trusted_numbers_in_attributable_months_are_summed():
    recs = [bill("Settembre 2025", 100),
            bill("Ottobre 2025", 999, kwh_status="UNCERTAIN"),            # not a trusted value
            bill("Novembre 2025", None),                           # no number at all
            bill("01/01/2025 - 31/03/2025", 300),                  # quarter: cannot be placed in one month
            bill("senza data", 400),                               # no period
            bill("31/07/2025 - 31/08/2025", 200)]                  # two months: split 100 + 100, flagged
    a = analyze_records("energy_utility", energy_schema(), recs)
    kwh = matrix_of(a, "consumo_kwh")
    assert kwh.total.totals == pytest.approx({2025: 300})
    assert kwh.total.values[2025][6] == kwh.total.values[2025][7] == 100        # Jul and Aug share the bimonthly bill
    assert (a.records_used, a.records_unplaced, a.records_spanning) == (2, 2, 1)


@pytest.mark.offline
def test_quantities_with_different_units_are_never_mixed():
    schema = resolve_schema("utility_consumption", "iso_14001")
    rec = lambda tipo, unit, qty, eur: {"fields": {   # noqa: E731
        "periodo_riferimento": _f("Settembre 2025"), "tipo_utenza": _f(tipo), "unita_misura": _f(unit),
        "consumo_totale": _f(qty), "totale_spesa_eur": _f(eur)}}
    a = analyze_records("iso_14001", schema, [rec("Luce", "kWh", 1000, 300), rec("Gas", "Smc", 200, 150)])
    qty = [m for m in a.matrices if m.kind == "quantity"]
    assert sorted(m.unit for m in qty) == ["Smc", "kWh"] and all(m.grand_total in (1000, 200) for m in qty)
    (cost,) = [m for m in a.matrices if m.kind == "cost"]
    assert cost.grand_total == 450 and [s.name for s in cost.by_dimension] == ["Gas", "Luce"]


@pytest.mark.offline
def test_documents_without_consumption_data_give_no_analysis():
    assert analyze_records("energy_utility", energy_schema(), []) is None
    assert analyze_records("energy_utility", energy_schema(), [bill("Settembre 2025", None)]) is None
    payload = {"iso_14001": {"regulatory_authorization": {"records": [{"fields": {"numero_atto_protocollo": _f("1")}}]},
                             "scadenziario": {"records": []}},
               "unknown_module": {"nope": {"records": [bill("Settembre 2025", 1)]}}}
    assert analyze_payload(payload) == [] and analyze_payload({}) == [] and analyze_payload(None) == []


@pytest.mark.offline
def test_analytics_block_is_validated():
    base = {"p": {"type": "string"}, "n": {"type": "number"}, "s": {"type": "string"}}
    good = {"period_fields": ["p"], "metrics": [{"field": "n", "kind": "cost"}]}
    assert _parse_schema("m", "s", {"fields": base, "analytics": good}, "t.json").analytics == good
    for bad in (
        {**good, "period_fields": []}, {**good, "period_fields": ["zzz"]}, {**good, "dimension_field": "zzz"},
        {**good, "max_span_months": 0}, {**good, "metrics": []},
        {**good, "metrics": [{"field": "s", "kind": "cost"}]},          # not a number field
        {**good, "metrics": [{"field": "n", "kind": "weight"}]},
        {**good, "metrics": [{"field": "n", "kind": "cost", "unit_field": "zzz"}]},
    ):
        with pytest.raises(ValueError):
            _parse_schema("m", "s", {"fields": base, "analytics": bad}, "t.json")


# ----------------------------------------------------------------------------- charts
@pytest.mark.offline
def test_chart_is_a_png_for_grouped_stacked_and_single_series():
    two_years = analyze_records("energy_utility", energy_schema(),
                                [bill("Settembre 2024", 800, 200), bill("Settembre 2025", 877, 230.5)])
    stacked = analyze_records("energy_utility", energy_schema(),
                              [bill("Settembre 2025", 100, pod="A"), bill("Settembre 2025", 50, pod="B")])
    single = analyze_records("energy_utility", energy_schema(), [bill("Settembre 2025", 100)])
    for analysis in (two_years, stacked, single):
        for m in analysis.matrices:
            png = render_monthly_chart(m, title="Titolo", dpi=80)
            data = png.getvalue()
            assert data.startswith(b"\x89PNG")
            assert Image.open(io.BytesIO(data)).size == (800, 368)


@pytest.mark.offline
def test_chart_returns_none_when_there_is_nothing_to_draw():
    a = analyze_records("energy_utility", energy_schema(), [bill("Settembre 2025", 1)])
    empty = type(a.matrices[0])(**{**a.matrices[0].__dict__, "years": ()})
    assert render_monthly_chart(empty) is None


# ----------------------------------------------------------------------------- Excel / Word integration
def _consumption_payload():
    return {"energy_utility": {"energy_bill": {"records": [
        bill("Settembre 2024", 800, 200), bill("Settembre 2025", 877, 230.5), bill("Ottobre 2025", 950, 260)]}}}


@pytest.mark.offline
def test_excel_gets_trend_sheet_with_matrix_and_embedded_charts():
    out = create_platform_excel_report([{"filename": "bolletta.pdf"}], _consumption_payload(),
                                       company_profile={"company_name": "X", "address": "Y"})
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert TREND_SHEET in wb.sheetnames
    ws = wb[TREND_SHEET]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    labels = {}
    for pos, v in sorted(cells.items()):                    # first occurrence = the consumption matrix (before the cost one)
        if isinstance(v, str):
            labels.setdefault(v, pos)
    r_tot, _ = labels["Totale"]
    assert [cells[(r_tot, 2)], cells[(r_tot, 3)]] == [800, 1827]            # year columns 2024, 2025
    r_sep, _ = labels["Settembre"]
    assert [cells[(r_sep, 2)], cells[(r_sep, 3)]] == [800, 877]
    r_dec = labels["Dicembre"][0]
    assert cells[(r_dec, 2)] == "-"                                          # no data stays '-', never 0
    assert "Var. % vs anno prec." in labels
    assert len(ws._images) == 2                                              # consumption + cost charts


@pytest.mark.offline
def test_excel_without_consumption_records_has_no_trend_sheet():
    payload = {"iso_14001": {"regulatory_authorization": {"records": [{
        "source_document": "a.pdf", "source_page": 1, "fields": {"numero_atto_protocollo": _f("1")}}]}}}
    out = create_platform_excel_report([{"filename": "a.pdf"}], payload)
    assert TREND_SHEET not in openpyxl.load_workbook(io.BytesIO(out)).sheetnames


@pytest.mark.offline
def test_word_report_gets_graphic_section_with_summary_table():
    data = {"energy_bill": {"records": _consumption_payload()["energy_utility"]["energy_bill"]["records"]}}
    doc = Document(io.BytesIO(create_styled_word_report(data, ["bolletta.pdf"], {"company_name": "X"},
                                                       module_id="energy_utility")))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(TREND_SECTION_TITLE in h for h in headings)
    assert len(doc.inline_shapes) == 2
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Spesa totale del periodo" in text and "Consumo totale" in text and "Mese di picco" in text
    assert "2.627,00 kWh" in text and "690,50 €" in text                  # 800 + 877 + 950 kWh; 200 + 230,5 + 260 €
    assert "Ottobre 2025 (950,00 kWh)" in text


@pytest.mark.offline
def test_word_report_without_consumption_records_has_no_graphic_section():
    data = {"regulatory_authorization": {"records": [{"source_page": 1, "fields": {
        "numero_atto_protocollo": _f("1")}}]}}
    doc = Document(io.BytesIO(create_styled_word_report(data, ["a.pdf"], {"company_name": "X"}, module_id="iso_14001")))
    assert not any(TREND_SECTION_TITLE in p.text for p in doc.paragraphs) and len(doc.inline_shapes) == 0


# ----------------------------------------------------------------------------- multi-month bills are apportioned
@pytest.mark.offline
@pytest.mark.parametrize("period,months", [
    ("01 MAGGIO 2026 - 30 GIUGNO 2026", (5, 6)),
    ("01 LUGLIO 2026 - 31 AGOSTO 2026", (7, 8)),
    ("01/05/2026 - 30/06/2026", (5, 6)),
    ("Maggio - Giugno 2026", (5, 6)),
])
def test_bimonthly_bill_is_split_equally_over_both_months(period, months):
    a = analyze_records("energy_utility", energy_schema(), [bill(period, 1000, 300.0)])
    kwh, cost = matrix_of(a, "consumo_kwh"), matrix_of(a, "costo_totale_eur")
    first, second = months
    for m, total in ((kwh, 1000), (cost, 300.0)):
        row = m.total.values[2026]
        assert row[first - 1] == row[second - 1] == pytest.approx(total / 2)        # 50% / 50%
        assert [v for i, v in enumerate(row) if i + 1 not in months and v is not None] == []   # no other month touched
        assert m.total.totals[2026] == pytest.approx(total)                         # nothing created or lost
        assert m.total.months_with_data[2026] == 2
    assert (a.records_used, a.records_spanning, a.records_unplaced) == (1, 1, 0)
    assert a.period_first == (2026, first) and a.period_last == (2026, second)
    assert kwh.n_records == 1                                                       # still one document


@pytest.mark.offline
def test_apportioning_keeps_monthly_bills_untouched_and_works_across_a_year_end_and_dimensions():
    recs = [bill("Aprile 2026", 400, pod="A"), bill("Dicembre 2025 - Gennaio 2026", 600, pod="A"),
            bill("Maggio 2026 - Giugno 2026", 101, pod="B")]
    kwh = matrix_of(analyze_records("energy_utility", energy_schema(), recs), "consumo_kwh")
    assert kwh.years == (2025, 2026)
    assert kwh.total.values[2025][11] == 300 and kwh.total.values[2026][0] == 300     # Dec 2025 / Jan 2026
    assert kwh.total.values[2026][3] == 400                                           # single month: whole value
    assert kwh.total.values[2026][4] == kwh.total.values[2026][5] == 50.5
    by_pod = {s.name: s for s in kwh.by_dimension}
    assert by_pod["A"].totals == pytest.approx({2025: 300, 2026: 700}) and by_pod["B"].totals == {2026: 101}


@pytest.mark.offline
def test_periods_longer_than_the_limit_are_still_not_apportioned():
    a = analyze_records("energy_utility", energy_schema(), [bill("Settembre 2025", 10), bill("01/01/2025 - 31/03/2025", 90)])
    assert a.records_unplaced == 1 and matrix_of(a, "consumo_kwh").grand_total == 10


@pytest.mark.offline
def test_excel_trend_sheet_states_that_values_are_apportioned():
    payload = {"energy_utility": {"energy_bill": {"records": [bill("01 LUGLIO 2026 - 31 AGOSTO 2026", 1000, 200.0)]}}}
    out = create_platform_excel_report([{"filename": "bolletta.pdf"}], payload)
    ws = openpyxl.load_workbook(io.BytesIO(out))[TREND_SHEET]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert text.count("valori ripartiti pro-quota sui mesi di competenza") >= 2        # header note + per-analysis note
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    rows = {}
    for (r, c), v in sorted(cells.items()):                  # first occurrence = the consumption matrix
        if c == 1 and isinstance(v, str):
            rows.setdefault(v, r)
    assert cells[(rows["Luglio"], 2)] == cells[(rows["Agosto"], 2)] == 500
    assert cells[(rows["Totale"], 2)] == 1000

