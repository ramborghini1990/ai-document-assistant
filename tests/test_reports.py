import io
import openpyxl
import pytest
from docx import Document
from app.reporting.excel_exporter import create_styled_excel_report
from app.reporting.word_exporter import create_styled_word_report


@pytest.mark.offline
def test_scoped_reporting_isolation():
    mixed_data = {
        "scadenziario": {
            "total_records": 1,
            "records": [{"source_document": "env.pdf", "source_page": 1, "fields": {"adempimento": {"normalized_value": "MUD", "status": "EXTRACTED"}}}]
        },
        "fuel_expense": {
            "total_records": 1,
            "records": [{"source_document": "fuel.xlsx", "source_page": 1, "fields": {"targa": {"normalized_value": "AA111BB", "status": "EXTRACTED"}}}]
        }
    }

    # ISO 14001 Excel isolation
    excel_bytes = create_styled_excel_report(mixed_data, ["env.pdf"], module_id="iso_14001")
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "SCADENZIARIO" in wb.sheetnames
    assert "FUEL_EXPENSE" not in wb.sheetnames

    # Fleet Word isolation
    word_bytes = create_styled_word_report(mixed_data, ["fuel.xlsx"], module_id="fleet_fuel_expenses")
    doc = Document(io.BytesIO(word_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "ISO 14001" not in full_text
    assert "FLOTTA" in full_text.upper()


@pytest.mark.offline
def test_word_report_regulatory_authorization_sections_and_empty_safety():
    data_with_prescriptions = {
        "regulatory_authorization": {
            "total_records": 1,
            "records": [{
                "source_document": "104-2531_2023.pdf",
                "source_page": 1,
                "fields": {
                    "numero_atto_protocollo": {"normalized_value": "104-2531", "status": "EXTRACTED"},
                    "ente_rilascio": {"normalized_value": "Provincia", "status": "EXTRACTED"},
                    "elenco_codici_eer": {"normalized_value": "15 01 01, 15 01 02", "status": "EXTRACTED"},
                    "quantitativi_movimentabili": {"normalized_value": "Fase 1: 3000 t/a", "status": "EXTRACTED"},
                    "quantitativi_deposito": {"normalized_value": "Fase 1: 150 t", "status": "EXTRACTED"},
                    "prescrizioni_rilevanti": {"normalized_value": "Garanzia finanziaria", "status": "EXTRACTED"},
                    "prescrizioni_con_scadenza": {
                        "normalized_value": "Trasmissione collaudo || entro 60 giorni\nRinnovo fideiussione || 31/12/2026",
                        "status": "EXTRACTED"
                    },
                    "prescrizioni_tecnico_gestionali": {
                        "normalized_value": "Mantenere pavimentazione impermeabile\nDivieto miscelazione rifiuti pericolosi",
                        "status": "EXTRACTED"
                    }
                }
            }]
        }
    }

    word_bytes = create_styled_word_report(
        data_with_prescriptions,
        ["104-2531_2023.pdf"],
        company_profile={"company_name": "inambiente S.r.l."},
        module_id="iso_14001"
    )
    doc = Document(io.BytesIO(word_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "RAPPORTO DI AUTORIZZAZIONE E PRESCRIZIONI AMBIENTALI — INAMBIENTE S.R.L." in full_text
    assert "3. Prescrizioni con Scadenza" in full_text
    assert "4. Prescrizioni Tecnico-Gestionali e Discorsive" in full_text
    assert "5. Approvazione e Visto" in full_text
    assert "1. Mantenere pavimentazione impermeabile" in full_text
    assert "2. Divieto miscelazione rifiuti pericolosi" in full_text

    # Table 0: 2-column vertical Key-Value specification block
    kv_table = doc.tables[0]
    assert len(kv_table.columns) == 2
    kv_pairs = {row.cells[0].text: row.cells[1].text for row in kv_table.rows}
    assert kv_pairs["N. Atto / Protocollo"] == "104-2531"
    assert kv_pairs["Ente Rilascio"] == "Provincia"
    assert "Attività e Categorie Autorizzate" in kv_pairs

    # Table 1: 3-column horizontal waste & quantities table (no Stato, no Prescrizioni)
    waste_table = doc.tables[1]
    waste_headers = [c.text for c in waste_table.rows[0].cells]
    assert waste_headers == [
        "Codici EER / Rifiuti Ammessi",
        "Quantitativi annui movimentabili",
        "Quantitativi autorizzati in deposito",
    ]
    assert "Stato" not in waste_headers
    assert "Prescrizioni / Condizioni Subordinanti" not in waste_headers
    assert waste_table.rows[1].cells[0].text == "15 01 01, 15 01 02"
    assert waste_table.rows[1].cells[1].text == "Fase 1: 3000 t/a"
    assert waste_table.rows[1].cells[2].text == "Fase 1: 150 t"

    # Table 2: 4-column Prescrizioni con Scadenza table
    presc_table = doc.tables[2]
    assert [c.text for c in presc_table.rows[0].cells] == [
        "N.", "Prescrizione / Condizione", "Scadenza / Termine", "Documento Fonte"
    ]
    assert presc_table.rows[1].cells[1].text == "Trasmissione collaudo"
    assert presc_table.rows[1].cells[2].text == "entro 60 giorni"
    assert presc_table.rows[2].cells[1].text == "Rinnovo fideiussione"
    assert presc_table.rows[2].cells[2].text == "31/12/2026"

    # Empty / None / "-" safety check
    data_empty_prescriptions = {
        "regulatory_authorization": {
            "total_records": 2,
            "records": [
                {
                    "source_document": "doc1.pdf",
                    "source_page": 1,
                    "fields": {
                        "numero_atto_protocollo": {"normalized_value": "A1", "status": "EXTRACTED"},
                        "prescrizioni_con_scadenza": {"normalized_value": None, "raw_value": "-", "status": "MISSING"},
                    }
                },
                {
                    "source_document": "doc2.pdf",
                    "source_page": 1,
                    "fields": {}
                }
            ]
        }
    }

    word_empty_bytes = create_styled_word_report(
        data_empty_prescriptions,
        ["doc1.pdf", "doc2.pdf"],
        module_id="iso_14001"
    )
    doc_empty = Document(io.BytesIO(word_empty_bytes))
    empty_text = "\n".join(p.text for p in doc_empty.paragraphs)
    assert empty_text.count("Nessuna prescrizione rilevata.") == 2
    assert "5. Approvazione e Visto" in empty_text
