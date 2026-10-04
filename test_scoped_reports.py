"""Offline test verifying strict module scoping in Excel and Word report deliverables."""
import io
import openpyxl
from docx import Document
from app.reporting.excel_exporter import create_styled_excel_report
from app.reporting.word_exporter import create_styled_word_report


def test_scoped_reports():
    print("Testing Step 3: Scoped Exporters & Cross-Module Isolation...")

    # داده ترکیبی شامل دو ماژول مختلف:
    # 1. scadenziario (متعلق به iso_14001)
    # 2. fuel_expense (متعلق به fleet_fuel_expenses)
    mixed_data = {
        "scadenziario": {
            "total_records": 1,
            "records": [
                {
                    "source_document": "audit_env.pdf",
                    "source_page": 1,
                    "fields": {
                        "tema": {"normalized_value": "Rifiuti", "status": "EXTRACTED"},
                        "adempimento": {"normalized_value": "MUD Annuale", "status": "EXTRACTED"},
                        "scadenza": {"normalized_value": "2026-04-30", "status": "EXTRACTED"},
                        "responsabile": {"normalized_value": "RSGA", "status": "EXTRACTED"}
                    }
                }
            ]
        },
        "fuel_expense": {
            "total_records": 1,
            "records": [
                {
                    "source_document": "fuel_card.xlsx",
                    "source_page": 1,
                    "fields": {
                        "targa": {"normalized_value": "AA123BB", "status": "EXTRACTED"},
                        "importo_eur": {"normalized_value": 120.5, "status": "EXTRACTED"}
                    }
                }
            ]
        }
    }

    # ۱. تولید اکسل با تعیین ماژول iso_14001
    excel_bytes = create_styled_excel_report(
        mixed_data,
        ["audit_env.pdf"],
        module_id="iso_14001"
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    
    assert "SCADENZIARIO" in sheet_names, "SCADENZIARIO sheet missing!"
    assert "FUEL_EXPENSE" not in sheet_names, "Contamination: FUEL_EXPENSE leaked into ISO 14001 report!"
    print(f"✓ Excel strictly isolated: Sheets={sheet_names} (No fuel_expense present).")

    # ۲. تولید ورد با تعیین ماژول fleet_fuel_expenses
    word_bytes = create_styled_word_report(
        mixed_data,
        ["fuel_card.xlsx"],
        module_id="fleet_fuel_expenses"
    )
    doc = Document(io.BytesIO(word_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    
    assert "ISO 14001" not in full_text, "Contamination: 'ISO 14001' found in Fleet report title!"
    assert "RAPPORTO CONSUMI E SPESE FLOTTA" in full_text, "Fleet custom title not found!"
    print("✓ Word report strictly isolated: Correct Fleet title and zero ISO 14001 mentions.")

    # ۳. بررسی آزمون فاز ۲۶ (سازگاری رو به عقب بدون تغییر)
    from test_phase26_configurable_export import test_phase26_configurable_export
    test_phase26_configurable_export()
    print("✓ Backward compatibility preserved: Phase 26 test passed seamlessly.")


if __name__ == "__main__":
    test_scoped_reports()
    print("\n🎉 STEP 3 (SCOPED EXPORTERS) FULLY VERIFIED & PASSED!")